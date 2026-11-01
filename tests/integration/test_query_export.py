"""``query_result(save_as=...)``: a query's whole result kept as a file for the user (CSV, or
Excel with the ``xlsx`` extra) through ``Runtime.save_file`` — offered only to a run with no
tool that runs code (``tools.export.runs_code``), which every adapter's tool set reads per run:
``ReAct`` and ``middleware()`` per model call, OpenAI Agents, Claude and MCP per attempt."""

from __future__ import annotations

import asyncio
import csv
import io
import json
from collections.abc import Iterator
from pathlib import Path
from types import SimpleNamespace
from typing import Any, Final

import pytest
from claude_agent_sdk import ClaudeAgentOptions
from langchain_core.messages import HumanMessage
from langchain_core.messages.utils import count_tokens_approximately
from langchain_core.utils.function_calling import convert_to_openai_tool
from openpyxl import load_workbook

from tests.support.adapters import BUILDERS, CLI
from tests.support.arm import arm
from tests.support.models import ScriptedChat
from tests.support.planned import FINAL, Call
from tests.support.sandbox import FakeSandboxes
from trellis import Harness, ReAct, sandbox, tool
from trellis.contracts import RunStatus
from trellis.harness import middleware as ours
from trellis.harness import runtime as runtime_module
from trellis.harness.files import FileStore
from trellis.harness.journal import Journal, content_key
from trellis.harness.tools import export, query, results
from trellis.harness.tools.convert import claude, mcp, openai_agents

#: 10,000 SKUs and their units in stock: 1,000 of them under 50 — more than a query's answer
#: shows (``query.ROWS``), every one of them in the file
SKUS: Final = [(f"SKU-{n:05d}", (n * 37) % 500) for n in range(1, 10_001)]
CSV: Final = "sku,stock\n" + "\n".join(f"{s},{u}" for s, u in SKUS)
KEPT: Final = f"/large_tool_results/{content_key('result', CSV)}"
LOW: Final = "SELECT sku, stock FROM result WHERE stock < 50 ORDER BY sku"
LOW_ROWS: Final = sorted((s, u) for s, u in SKUS if u < 50)


@tool(side_effects="read")
def inventory() -> str:
    """Every SKU with its units in stock, as CSV."""
    return CSV


@pytest.fixture
def stored(harness: Harness, tmp_path: Path) -> Path:
    """The harness's files kept below a folder of the test's own."""
    root = tmp_path / "files"
    harness.files = FileStore(root.as_uri())
    return root


def kept(root: Path, name: str) -> Path:
    [found] = list(root.rglob(name))
    return found


def saved(root: Path) -> list[Path]:
    return [p for p in root.rglob("*") if p.is_file()]


def sheet_rows(path: Path) -> list[tuple[Any, ...]]:
    book = load_workbook(path, read_only=True)
    return [tuple(r) for r in book.worksheets[0].iter_rows(values_only=True)]


def save_plan(**save: Any) -> list[Call]:
    return [("inventory", {}), ("query_result", {"file_path": KEPT, "sql": LOW, **save})]


# --------------------------------------------------------------------------- every adapter
@pytest.mark.parametrize("framework", ["react", "middleware", "openai_agents"])
async def test_every_framework_saves_a_querys_whole_result_as_excel(
    harness: Harness, engine: str, stored: Path, framework: str, tmp_path: Path
) -> None:
    plan = save_plan(save_as="xlsx", file_name="low-stock")
    target, own = await BUILDERS[framework](harness, [inventory], tmp_path, plan)
    result = await arm(harness.wrap(target, id=f"save-{framework}"), own).run("q", user="u")
    assert result.status is RunStatus.SUCCESS, result.error
    assert result.answer.startswith(
        f"Done. Saved low-stock.xlsx for the user: {len(LOW_ROWS)} rows"
    ), result.answer
    rows = sheet_rows(kept(stored, "low-stock.xlsx"))
    assert rows == [("sku", "stock"), *LOW_ROWS]  # every row, exactly: not the 200 shown


async def test_claude_without_bash_saves_and_with_bash_writes_with_its_code(
    harness: Harness, stored: Path, tmp_path: Path
) -> None:
    """Claude Code's ``Bash`` runs code: with it, ``query_result`` has no ``save_as`` (the call
    is refused, naming the argument); without it, the file is saved."""
    script = [
        {"tool": "inventory", "args": {}},
        {"tool": "query_result", "args": {"file_path": KEPT, "sql": LOW, "save_as": "csv"}},
        {"text": FINAL},
    ]

    def options(**more: Any) -> ClaudeAgentOptions:
        return ClaudeAgentOptions(
            cli_path=CLI, system_prompt="s", env={"FAKE_CLAUDE_SCRIPT": json.dumps(script)}, **more
        )

    saving = arm(harness.wrap(options(disallowed_tools=["Bash"]), id="claude-save"), [inventory])
    result = await saving.run("q", user="u")
    assert result.status is RunStatus.SUCCESS, result.error
    assert "Saved query_result.csv for the user" in result.answer, result.answer
    text = kept(stored, "query_result.csv").read_text()
    assert list(csv.reader(io.StringIO(text)))[1:] == [[s, str(u)] for s, u in LOW_ROWS]

    coding = arm(harness.wrap(options(), id="claude-bash"), [inventory])
    result = await coding.run("q", user="u")
    assert "'save_as' was unexpected" in result.answer, result.answer  # the SDK's check


async def test_a_run_with_a_sandbox_writes_its_files_with_code_not_save_as(
    harness: Harness, stored: Path
) -> None:
    asked = {"file_path": KEPT, "sql": LOW, "save_as": "csv"}
    model = ScriptedChat([("inventory", {}), ("query_result", asked), "done"])
    tools = [inventory, sandbox(FakeSandboxes()).tools()]
    agent = arm(harness.wrap(ReAct(system_prompt="s", model=model), id="sbx-save"), tools)
    assert (await agent.run("q", user="u")).status is RunStatus.SUCCESS
    offered = {t["function"]["name"]: t["function"] for t in model.requests[1]["tools"]}
    assert "save_as" not in offered["query_result"]["parameters"]["properties"]
    told = str(model.requests[2]["messages"][-1]["content"])
    assert told.startswith("query_result was not run: unknown argument(s): save_as"), told
    assert not saved(stored)


# --------------------------------------------------------------------------- the variants
class Run:
    """A run in progress, as the converters read it: its framework's tools, and what it
    keeps."""

    def __init__(self, *, native: tuple[str, ...] = ()) -> None:
        self.replay = SimpleNamespace(journal=Journal())
        self.toolbox: dict[str, Any] = {}
        self.native = dict.fromkeys(native)


@pytest.fixture(params=[(), ("sandbox_exec",), ("Bash",)], ids=["no-code", "sandbox", "bash"])
def run(request: pytest.FixtureRequest) -> Iterator[tuple[Run, bool]]:
    """A run with no code tool (it saves), with the harness's sandbox, with Claude's Bash."""
    names: tuple[str, ...] = request.param
    given = Run(native=names)  # the toolbox's own: a sandbox's tools (above)
    token = runtime_module._current.set(given)  # type: ignore[arg-type]
    yield given, not names
    runtime_module._current.reset(token)


def saves(schema: dict[str, Any]) -> bool:
    return "save_as" in schema["properties"]


async def test_each_converter_offers_the_runs_own_variant(run: tuple[Run, bool]) -> None:
    _, saving = run
    [function] = [t for t in openai_agents.convert([]) if t.name == "query_result"]
    assert saves(function.params_json_schema) is saving
    for server in (mcp.convert([]), claude.convert([])["instance"]):
        handler: Any = server._request_handlers["tools/list"].handler
        listed = await handler(None, None)
        [found] = [t for t in listed.tools if t.name == "query_result"]
        assert saves(found.input_schema) is saving
    assert results.readers() is (results.SAVING if saving else results.READERS)


async def test_react_offers_the_variant_per_model_call(run: tuple[Run, bool]) -> None:
    given, saving = run
    given.replay.journal.results["/large_tool_results/x"] = "a,b\n1,2"
    sent: list[Any] = []

    async def handler(request: Any) -> Any:
        sent.extend(request.tools)

    request = SimpleNamespace(
        tools=[],
        messages=[],
        state={},
        override=lambda **c: SimpleNamespace(**{"state": {}, "messages": [], **c}),
    )
    await ours.HarnessTools().awrap_model_call(request, handler)  # type: ignore[arg-type]
    [found] = [t for t in sent if t.name == "query_result"]
    assert found is (ours.SAVING if saving else ours.READERS)["query_result"]


def test_saving_adds_a_few_dozen_tokens_to_query_result_only() -> None:
    def tokens(found: Any) -> int:
        return count_tokens_approximately(
            [HumanMessage("x")], tools=[convert_to_openai_tool(found)]
        )

    added = tokens(ours.SAVING["query_result"]) - tokens(ours.READERS["query_result"])
    assert 0 < added <= 60, added
    assert ours.SAVING["read_file"] is ours.READERS["read_file"]  # the same tool, as it was


# --------------------------------------------------------------------------- the file
@pytest.fixture
def root(tmp_path: Path, engine: str) -> Iterator[Path]:
    """A run saving files (no code tool) in a store at the folder yielded, on each engine."""
    found = tmp_path / "files"
    store = FileStore(found.as_uri())

    async def save_file(name: str, data: bytes, *, source: str) -> Any:
        return await store.put("f", name, data, source=source)

    given = Run()
    given.save_file = save_file  # type: ignore[attr-defined]
    token = runtime_module._current.set(given)  # type: ignore[arg-type]
    yield found
    runtime_module._current.reset(token)


def ask(**args: Any) -> str:
    path = results.keep("n,name,blob\n1,a,\n2,b,x\n3,c,y")
    assert path is not None
    return asyncio.run(query.query({"file_path": path, **args}))


def test_a_csv_has_every_row_and_its_answer_the_first_few(root: Path) -> None:
    said = ask(sql="SELECT n, name FROM result ORDER BY n", save_as="csv", file_name="all.csv")
    assert said.startswith("Saved all.csv for the user: 3 rows, 19 bytes;"), said
    assert said.endswith("Its first rows:\nn,name\n1,a\n2,b\n3,c")
    assert kept(root, "all.csv").read_text() == "n,name\n1,a\n2,b\n3,c\n"


def test_an_excel_cell_that_is_no_number_or_text_is_its_text(root: Path) -> None:
    said = ask(sql="SELECT n, NULL AS z, CAST(name AS BLOB) AS b FROM result", save_as="xlsx")
    assert said.startswith("Saved query_result.xlsx for the user: 3 rows"), said
    rows = sheet_rows(kept(root, "query_result.xlsx"))
    assert rows == [("n", "z", "b"), (1, None, "b'a'"), (2, None, "b'b'"), (3, None, "b'c'")]


def test_a_saved_file_keeps_the_engines_values(root: Path, engine: str) -> None:
    """What DuckDB answers with that SQLite never does — a decimal, a boolean — is a number and
    a boolean in Excel, and exact text, ``true``/``false`` in CSV, as the answer shows it."""
    sql = "SELECT n, n * 1.5 AS d, n > 1 AS b FROM result ORDER BY n"
    said = ask(sql=sql, save_as="xlsx", file_name="v")
    flags = {"duckdb": (False, True, True), "sqlite": (0, 1, 1)}[engine]
    rows = sheet_rows(kept(root, "v.xlsx"))
    assert rows == [("n", "d", "b"), *zip((1, 2, 3), (1.5, 3.0, 4.5), flags, strict=True)]
    shown = {"duckdb": ("false", "true"), "sqlite": ("0", "1")}[engine]
    assert said.endswith(f"n,d,b\n1,1.5,{shown[0]}\n2,3.0,{shown[1]}\n3,4.5,{shown[1]}"), said
    ask(sql=sql, save_as="csv", file_name="v")
    text = kept(root, "v.csv").read_text().splitlines()
    exact = {"duckdb": ["1.5", "3.0", "4.5"], "sqlite": ["1.5", "3.0", "4.5"]}[engine]
    assert [r.split(",")[1:] for r in text[1:]] == [
        [d, f] for d, f in zip(exact, (shown[0], shown[1], shown[1]), strict=True)
    ]


@pytest.mark.parametrize(
    ("args", "said"),
    [
        ({"save_as": "pdf"}, "save_as is one of: csv, xlsx"),
        ({"save_as": "csv", "file_name": "../up"}, "'../up.csv' is not a file name"),
    ],
)
def test_a_save_the_harness_cannot_make_is_not_run(
    root: Path, args: dict[str, Any], said: str
) -> None:
    told = ask(sql="SELECT 1", **args)
    assert told.startswith(f"query_result was not run: {said}"), told


def test_more_rows_than_a_file_holds_are_refused(
    root: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(export, "MAX_ROWS", 2)
    for kind in export.formats():
        told = ask(sql="SELECT n FROM result", save_as=kind)
        assert told == (
            "Error: the query has more than 2 rows: a file holds at most that many: "
            "aggregate, or narrow it with WHERE"
        ), told
    assert not root.exists()


def test_a_failed_query_saves_nothing(root: Path) -> None:
    told = ask(sql="SELECT nope FROM result", save_as="csv")
    assert told.startswith("Error: the query failed:") and "nope" in told, told
    assert not root.exists()


def test_excel_is_offered_only_with_openpyxl(monkeypatch: pytest.MonkeyPatch) -> None:
    export.formats.cache_clear()
    monkeypatch.setattr(export.importlib.util, "find_spec", lambda name: None)
    try:
        assert export.formats() == ("csv",)
    finally:
        monkeypatch.undo()
        export.formats.cache_clear()
    assert export.formats() == ("csv", "xlsx")
