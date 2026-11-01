"""The files kept for the user (``harness.files``: an export, a sandbox's output) read back in a
later run — listed, read as rows, queried beside the run's own results — always within the
run's tenant and user: a path names a conversation and a run of theirs, never anyone else's.

The model reads them at ``/files/<run>/<name>`` (a file of this chat), at
``/files/<conversation>/<run>/<name>`` (another chat's), and ``/files`` lists every chat's. A
run of a chat that has files is told of them as it starts (:func:`note`: a line each, with
what made it and when the results it was made from were fetched, from the file's record:
``FileStore.record``); a run of a user with files only in other chats is told where to list
them. ``read_file`` reads one — CSV and Excel as CSV rows (Excel with the ``xlsx`` extra) —
its first page saying how it was made: the run, the query, its sources and their fetch times,
the run's tool calls. ``query_result`` loads one from the store as a table (:func:`table`), so
one query compares an earlier file with this run's results.
"""

from __future__ import annotations

import asyncio
import csv
import io
import zipfile
from dataclasses import dataclass
from datetime import datetime
from typing import Any, Final

from trellis.contracts import ToolError
from trellis.harness import snapshots
from trellis.harness.files import FileStore, file_name, part
from trellis.harness.runtime import Runtime, current
from trellis.harness.tools import results

#: Where the model reads the files kept for its user.
PREFIX: Final = "/files"
#: The most files a run is told of as it starts, and listed at once (``read_file("/files")``);
#: and the most a listing reads the records of.
NOTED: Final = 20
LISTED: Final = 100
#: The largest file read back (a query loads at most as much: ``query.LOAD_CHARS``).
MAX_BYTES: Final = 64 * 1024 * 1024
NOTE: Final = (
    "Files kept for the user in this chat (it is now {now}). read_file(file_path=...) reads "
    "one: its rows, and "
    "how it was made. query_result(file_path=..., sql=...) queries one as the table result, "
    "joined with this run's results by their tools' names. read_file(file_path=\"/files\") "
    "lists the user's files of every chat."
)
ELSEWHERE: Final = (
    'The user has files kept in other chats: read_file(file_path="/files") lists them.'
)
MISSING: Final = (
    "Error: no file is kept at {path}: it was deleted, or never kept. "
    'read_file(file_path="/files") lists the user\'s files.'
)
GONE: Final = "Error: {path} was deleted: it is no longer kept. It was {made}."
#: The extension Excel's files have.
XLSX: Final = ".xlsx"

#: A table read from a file: its columns and its rows.
Table = tuple[list[str], list[list[Any]]]


@dataclass(frozen=True, slots=True)
class Kept:
    """A file kept for the user: the path the model reads it at, where it is in the store, its
    size, and its record (how it was made: ``FileStore.record``)."""

    path: str
    folder: str
    name: str
    #: ``None``: the file was deleted (its record is left)
    size: int | None
    record: dict[str, Any]

    def line(self) -> str:
        """The file in one line: its path, size, and how it was made (:func:`made`) — or that
        it was deleted."""
        if self.size is None:
            return f"- {self.path}: deleted, no longer kept ({made(self.record)})"
        return f"- {self.path} ({self.size} bytes; {made(self.record)})"


def made(record: dict[str, Any]) -> str:
    """How a kept file was made, from its record, in a few words: when and by what it was
    saved, and from what results fetched when."""
    said = f"saved {when(record.get('saved_at'))} by {record.get('source') or 'an unknown tool'}"
    sources = ", ".join(
        f"{s.get('tool') or s.get('path')} fetched {when(s.get('fetched_at'))}"
        for s in record.get("from") or ()
    )
    return f"{said} from {sources}" if sources else said


def is_stored(path: str) -> bool:
    """Whether ``path`` (as the model gave it, with a leading ``/``) is a kept file's or the
    listing's."""
    return path == PREFIX or path.startswith(f"{PREFIX}/")


def when(text: Any) -> str:
    """An ISO time as the model reads it (``2026-10-31 10:02 UTC``); ``an unknown time``."""
    try:
        found = datetime.fromisoformat(str(text))
    except ValueError:
        return "an unknown time"
    return found.strftime("%Y-%m-%d %H:%M UTC")


def _user(runtime: Runtime) -> str:
    return f"{part(runtime.tenant)}/{part(runtime.user)}"


def where(runtime: Runtime, path: str) -> tuple[str, str] | None:
    """The folder and name in the store of the file the model reads at ``path`` — within the
    run's tenant and user, whatever the path says — or ``None``: no kept file's path."""
    parts = path[len(PREFIX) + 1 :].split("/")
    if len(parts) == 2:  # <run>/<name>: this chat's
        parts.insert(0, part(runtime.conversation))
    if len(parts) != 3 or any(p != part(p) for p in parts[:2]):
        return None
    try:
        name = file_name(parts[2])
    except ToolError:
        return None
    return f"{_user(runtime)}/{parts[0]}/{parts[1]}", name


async def kept(runtime: Runtime, *, everywhere: bool, store: FileStore | None = None) -> list[Kept]:
    """The files kept for the run's user — in its conversation, or ``everywhere`` — newest
    first, at most :data:`LISTED`, each with its record (read at once)."""
    store = store or runtime.agent.harness.files
    here = part(runtime.conversation)
    base = _user(runtime) if everywhere else f"{_user(runtime)}/{here}"
    found = sorted(await store.files(base))[-LISTED:]
    records = await asyncio.gather(*(store.record(where, name) for where, name, _ in found))
    listed = []
    for (folder, name, size), record in zip(found, records, strict=True):
        conversation, _, run = folder.split("/", 2)[2].partition("/")
        chat = "" if conversation == here else f"{conversation}/"
        path = f"{PREFIX}/{chat}{run}/{name}"
        listed.append(Kept(path, folder, name, size, record or {}))
    listed.sort(key=lambda k: str(k.record.get("saved_at") or ""), reverse=True)
    return listed


def looks(runtime: Runtime) -> bool:
    """Whether a run is told of the files kept for its user (:func:`note`): not a sub-agent's
    (its parent was told), nor a function's (its code is offered no reader: ``agent.tools`` is
    its toolbox)."""
    return runtime.parent is None and runtime.agent.adapter.name != "function"


async def note(runtime: Runtime) -> str | None:
    """What a run is told as it starts of the files kept for its user: its chat's (:data:`NOTE`,
    the :data:`NOTED` newest, a line each), else that other chats have some (:data:`ELSEWHERE`:
    a folder of another chat, one listing); ``None`` when there are none, or when the store
    cannot say (a warning). Asked only of a run that :func:`looks`."""
    store = runtime.agent.harness.files
    try:
        here = (
            []
            if runtime.conversation == runtime.run_id
            else await kept(runtime, everywhere=False, store=store)
        )
        if here:
            more = len(here) - NOTED
            lines = [k.line() for k in here[:NOTED]]
            if more > 0:
                lines.append(f'- and {more} more: read_file(file_path="/files")')
            return "\n".join([NOTE.format(now=when(snapshots.now().isoformat())), *lines])
        chats = await store.folders(_user(runtime))
        return ELSEWHERE if set(chats) - {part(runtime.conversation)} else None
    except Exception as exc:
        runtime.events.warning("files", f"the files kept for the user: {exc}")
        return None


def told(runtime: Runtime, said: str | None) -> None:
    """The run told ``said`` (:func:`note`), after its memory context: ``read_file`` and
    ``query_result`` are offered from its first model call (``results.kept``)."""
    if said is None:
        return
    runtime.files_note = said
    runtime.context = "\n\n".join(p for p in (runtime.pushed_context, said) if p)


async def read(path: str, offset: int = 0, limit: int = results.READ_LINES) -> str:
    """``read_file`` of a kept file: ``/files`` lists the user's files of every chat; a file's
    lines from ``offset`` (CSV and Excel as CSV rows), its first page preceded by how it was
    made (:func:`how`)."""
    runtime = current()
    if runtime is None:
        return f"Error: File '{path}' not found"
    if path.rstrip("/") == PREFIX:
        listed = await kept(runtime, everywhere=True)
        if not listed:
            return "No files are kept for the user."
        return "\n".join(["The user's files, newest first:", *(k.line() for k in listed)])
    found = await _text(path)
    if isinstance(found, str):
        return found
    text, record = found
    page = results.page(text, path, offset, limit)
    return f"{how(path, record)}\n\n{page}" if offset <= 0 else page


async def table(path: str) -> tuple[Table | str, dict[str, Any]] | str:
    """A kept file as ``query_result`` loads it — a table (Excel's rows), or its text (CSV,
    JSON, lines: ``query.rows_of`` reads it) — and its record; or what is wrong."""
    runtime = current()
    if runtime is None:
        return f"Error: File '{path}' not found"
    found = await _loaded(runtime, path)
    if isinstance(found, str):
        return found
    data, name, record = found
    try:
        if name.lower().endswith(XLSX):  # parsed in a thread: the event loop never waits on it
            return await asyncio.to_thread(_sheet, data), record
        return data.decode(), record
    except (ValueError, KeyError, ImportError, OSError, zipfile.BadZipFile) as exc:
        return f"Error: {path} is no table: {type(exc).__name__}: {exc}"


async def _text(path: str) -> tuple[str, dict[str, Any]] | str:
    """A kept file as text: CSV rows for an Excel file."""
    found = await table(path)
    if isinstance(found, str):
        return found
    content, record = found
    if isinstance(content, str):
        return content, record
    columns, rows = content
    out = io.StringIO()
    writer = csv.writer(out, lineterminator="\n")
    writer.writerow(columns)
    writer.writerows(rows)
    return out.getvalue(), record


async def _loaded(runtime: Runtime, path: str) -> tuple[bytes, str, dict[str, Any]] | str:
    located = where(runtime, path)
    if located is None:
        return (
            f"Error: {path} is no kept file's path (/files/<run>/<name>, or "
            '/files/<chat>/<run>/<name>): read_file(file_path="/files") lists them'
        )
    folder, name = located
    store = runtime.agent.harness.files
    size = await store.size(folder, name)
    if size is None:
        record = await store.record(folder, name)
        if record is None:
            return MISSING.format(path=path)
        return GONE.format(path=path, made=made(record))
    if size > MAX_BYTES:
        return f"Error: {path} is {size} bytes, more than {MAX_BYTES} read here"
    data, record = await asyncio.gather(store.read(folder, name), store.record(folder, name))
    if data is None:  # deleted meanwhile
        return MISSING.format(path=path)
    return data, name, record or {}


def _sheet(data: bytes) -> Table:
    """An Excel file's first sheet: its first row the columns, the rest its rows (a time as
    its ISO text)."""
    from openpyxl import load_workbook  # noqa: PLC0415 - the xlsx extra

    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        rows = book.worksheets[0].iter_rows(values_only=True)
        head = next(rows, ())
        columns = [str(c) if c is not None else f"c{n + 1}" for n, c in enumerate(head)]
        width = len(columns)
        found = [
            [v.isoformat() if isinstance(v, datetime) else v for v in row][:width]
            for row in rows
            if any(v is not None for v in row)
        ]
    finally:
        book.close()
    return columns, found


def how(path: str, record: dict[str, Any]) -> str:
    """How the file at ``path`` was made, from its record: when, in which run, by what — the
    query, the results it read and when each was fetched — and the run's tool calls."""
    if not record:
        return f"{path}: how it was made is not recorded."
    said = (
        f"{path} was saved {when(record.get('saved_at'))} in run {record.get('run')} by "
        f"{record.get('source')}"
    )
    if record.get("sql"):
        said += f" with the query {record['sql']}"
    sources = ", ".join(
        f"{s.get('table')} = {s.get('tool') or s.get('path')} fetched {when(s.get('fetched_at'))}"
        for s in record.get("from") or ()
    )
    if sources:
        said += f", over {sources}"
    calls = "; ".join(
        f"{c.get('tool')}({_args(c.get('args'))}) {c.get('status')}"
        for c in record.get("calls") or ()
    )
    return f"{said}. The run's tool calls: {calls or 'none'}."


def _args(args: Any) -> str:
    text = ", ".join(f"{k}={v!r}" for k, v in (args or {}).items())
    return text if len(text) <= ARGS_CHARS else text[:ARGS_CHARS] + "…"


#: The most of a call's arguments :func:`how` shows.
ARGS_CHARS: Final = 300


__all__ = [
    "PREFIX",
    "Kept",
    "is_stored",
    "kept",
    "looks",
    "note",
    "read",
    "table",
    "told",
    "where",
]
